"""中国輸出Invoice管理: パッキングリスト形式Excelの判定・パース・INVOICE_NO集約。

基幹システムから出力されるパッキングリストExcel（1行=1品目明細）を読み、
INVOICE_NOごとに集約したInvoice実績の元データを返す。列は1シート目1行目の
ヘッダー名で特定するため、列の順序や余分な列の有無は問わない。
"""
import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import openpyxl

REQUIRED_COLUMNS = ('INVOICE_NO', '出荷日', '金額_取引')
CURRENCY_COLUMN = '通貨CD'
_DATE_FORMATS = ('%Y-%m-%d', '%Y/%m/%d')


class PackingListParseError(Exception):
    """パース失敗。errorsに行番号付きメッセージのリストを持つ。"""

    def __init__(self, errors):
        self.errors = list(errors)
        super().__init__('; '.join(self.errors))


def _load_first_sheet(uploaded_file):
    uploaded_file.seek(0)
    wb = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
    return wb, wb.worksheets[0]


def _header_map(ws):
    """1行目のヘッダー名 -> 0始まり列index。"""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return {
            str(v).strip(): idx
            for idx, v in enumerate(row)
            if v is not None and str(v).strip()
        }
    return {}


def is_packing_list_file(uploaded_file):
    """必須列のヘッダーが揃った.xlsxならTrue。読めないファイルはFalse（従来どおり添付扱いに落とす）。"""
    try:
        wb, ws = _load_first_sheet(uploaded_file)
    except Exception:
        uploaded_file.seek(0)
        return False
    try:
        headers = _header_map(ws)
    finally:
        if wb is not None:
            wb.close()
        uploaded_file.seek(0)
    return all(col in headers for col in REQUIRED_COLUMNS)


def _parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(text)


def parse_packing_list(uploaded_file):
    """INVOICE_NOごとに集約した [{'invoice_no', 'invoice_total': Decimal, 'export_date': date}] を返す。

    - invoice_total = 金額_取引の合計（小数2桁・四捨五入）
    - export_date = 出荷日（同一Invoice内で割れていた場合は最大値）
    - INVOICE_NOが空の行・全列空の行は読み飛ばす
    - 値が解釈できない行・通貨混在・集約結果0件は PackingListParseError（部分取り込みはしない）
    """
    wb = None
    try:
        wb, ws = _load_first_sheet(uploaded_file)
        headers = _header_map(ws)
        missing = [c for c in REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise PackingListParseError([f'必須列がありません: {"、".join(missing)}'])
        currency_idx = headers.get(CURRENCY_COLUMN)

        errors = []
        groups = {}  # invoice_no -> {'total': Decimal, 'dates': set, 'currencies': set}
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None or str(v).strip() == '' for v in row):
                continue

            def cell(name):
                idx = headers[name]
                return row[idx] if idx < len(row) else None

            invoice_no = str(cell('INVOICE_NO') or '').strip()
            if not invoice_no:
                continue
            group = groups.setdefault(
                invoice_no, {'total': Decimal('0'), 'dates': set(), 'currencies': set()})

            amount_raw = cell('金額_取引')
            try:
                group['total'] += Decimal(str(amount_raw).strip())
            except InvalidOperation:
                errors.append(f'{row_no}行目: 金額_取引を数値として読み取れません（{amount_raw}）')

            date_raw = cell('出荷日')
            try:
                group['dates'].add(_parse_date(date_raw))
            except (TypeError, ValueError):
                errors.append(f'{row_no}行目: 出荷日を日付として読み取れません（{date_raw}）')

            if currency_idx is not None:
                raw = row[currency_idx] if currency_idx < len(row) else None
                currency = str(raw or '').strip()
                if currency:
                    group['currencies'].add(currency)

        for invoice_no, group in groups.items():
            if len(group['currencies']) > 1:
                errors.append(
                    f'{invoice_no}: 通貨CDが混在しています（{"、".join(sorted(group["currencies"]))}）')
        if not groups and not errors:
            errors.append('INVOICE_NOが入力された行がありません。')
        if errors:
            raise PackingListParseError(errors)

        return [
            {
                'invoice_no': invoice_no,
                'invoice_total': group['total'].quantize(
                    Decimal('0.01'), rounding=ROUND_HALF_UP),
                'export_date': max(group['dates']),
            }
            for invoice_no, group in sorted(groups.items())
        ]
    finally:
        if wb is not None:
            wb.close()
        uploaded_file.seek(0)
