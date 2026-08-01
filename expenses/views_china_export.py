"""各部報告: 中国輸出実績報告 (T_ChinaExport) の一覧・入力・Excel出力ビュー"""
import logging
from urllib.parse import urlencode

from datetime import date, datetime
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.contrib import messages
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from .forms import ChinaExportUpdateForm
from .models import T_ChinaExport
from .exchange_upload import get_field_mapping, parse_excel_rows

logger = logging.getLogger(__name__)

_SORT_FIELDS = {
    'order_no', 'supplier_cd', 'supplier_name', 'item_name1', 'item_name2',
    'unit_price', 'purchase_date', 'quantity', 'amount', 'account_cd', 'account_name',
    'burden_bumon_cd', 'burden_bumon_name', 'order_bumon_name', 'order_staff_name',
    'export_planned_date', 'export_date', 'invoice_no',
}
_DEFAULT_SORT = 'purchase_date'


def _require_china_export_access(user):
    if not (user.has_role('export') or user.has_role('admin')):
        raise PermissionDenied()


def _resolve_sort(request):
    """GETの`sort`パラメータを検証する。'-field'は降順。不正/未知の値はデフォルトにフォールバック。"""
    raw = request.GET.get('sort', _DEFAULT_SORT)
    field = raw[1:] if raw.startswith('-') else raw
    if field not in _SORT_FIELDS:
        return _DEFAULT_SORT
    return raw


def _china_export_queryset(show_all, sort=_DEFAULT_SORT):
    records = T_ChinaExport.objects.order_by(sort, 'pk')
    if not show_all:
        records = records.filter(export_date__isnull=True)
    return records


def _build_sort_links(show_all, sort_raw):
    """各列見出し用のソートリンク情報を組み立てる。"""
    links = {}
    for field in _SORT_FIELDS:
        next_sort = f'-{field}' if sort_raw == field else field
        params = {'sort': next_sort}
        if show_all:
            params['show'] = 'all'
        if sort_raw == field:
            arrow = '▲'
        elif sort_raw == f'-{field}':
            arrow = '▼'
        else:
            arrow = ''
        links[field] = {
            'url': f'?{urlencode(params)}',
            'arrow': arrow,
            'active': arrow != '',
        }
    return links


@login_required
def china_export_list(request):
    _require_china_export_access(request.user)
    show_all = request.GET.get('show') == 'all'
    sort_raw = _resolve_sort(request)
    return render(request, 'expenses/china_export_list.html', {
        'records': _china_export_queryset(show_all, sort_raw),
        'show_all': show_all,
        'sort_links': _build_sort_links(show_all, sort_raw),
        'current': 'china_export_list',
    })


_EXCEL_HEADERS = [
    '状態', '注文番号', '仕入先コード', '仕入先名', '品目名1', '品目名2',
    '仕入単価', '購入日', '数量', '金額', '科目コード', '科目名',
    '負担部門コード', '負担部署名', '発注部署名', '発注担当名',
    '輸出予定日', '輸出日', 'インボイスNo',
]
_EXCEL_COLUMN_WIDTHS = [8, 14, 12, 20, 24, 24, 12, 12, 10, 12, 10, 16, 12, 14, 14, 14, 12, 12, 16]
_EXCEL_DATE_COLS = {8, 17, 18}    # 購入日, 輸出予定日, 輸出日
_EXCEL_MONEY_COLS = {7, 10}       # 仕入単価, 金額
_EXCEL_QTY_COL = 9                # 数量


def _china_export_to_excel_row(r):
    return [
        '輸出済' if r.export_date else '未輸出',
        r.order_no or '',
        r.supplier_cd or '',
        r.supplier_name or '',
        r.item_name1,
        r.item_name2 or '',
        float(r.unit_price) if r.unit_price is not None else None,
        r.purchase_date,
        float(r.quantity) if r.quantity is not None else None,
        float(r.amount),
        r.account_cd or '',
        r.account_name or '',
        r.burden_bumon_cd or '',
        r.burden_bumon_name or '',
        r.order_bumon_name or '',
        r.order_staff_name or '',
        r.export_planned_date,
        r.export_date,
        r.invoice_no or '',
    ]


@login_required
def china_export_excel(request):
    """一覧を成形済みのExcel(.xlsx)としてダウンロードする（表示中の未輸出のみ/全件の状態を引き継ぐ）。"""
    _require_china_export_access(request.user)
    show_all = request.GET.get('show') == 'all'
    records = _china_export_queryset(show_all)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '中国輸出実績報告'

    ws.append(_EXCEL_HEADERS)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='495057', end_color='495057', fill_type='solid')
    for col_idx in range(1, len(_EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r in records.iterator():
        ws.append(_china_export_to_excel_row(r))

    last_row = ws.max_row
    for row_idx in range(2, last_row + 1):
        for col_idx in _EXCEL_DATE_COLS:
            ws.cell(row=row_idx, column=col_idx).number_format = 'yyyy/mm/dd'
        for col_idx in _EXCEL_MONEY_COLS:
            ws.cell(row=row_idx, column=col_idx).number_format = '#,##0'
        ws.cell(row=row_idx, column=_EXCEL_QTY_COL).number_format = '#,##0.00'

    for col_idx, width in enumerate(_EXCEL_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="china_export.xlsx"'
    wb.save(response)
    return response


@login_required
@require_POST
def china_export_bulk_update(request):
    """一覧に表示中の全行分の輸出予定日/輸出日/インボイスNoを一括保存する。
    各入力欄は `<field>_<pk>` という名前で送信され、`pks` (複数値) が対象レコードの一覧を表す。"""
    _require_china_export_access(request.user)
    pks = request.POST.getlist('pks')
    records = T_ChinaExport.objects.filter(pk__in=pks)
    for record in records:
        data = {
            'export_planned_date': request.POST.get(f'export_planned_date_{record.pk}', ''),
            'export_date': request.POST.get(f'export_date_{record.pk}', ''),
            'invoice_no': request.POST.get(f'invoice_no_{record.pk}', ''),
        }
        form = ChinaExportUpdateForm(data, instance=record)
        if form.is_valid():
            updated = form.save(commit=False)
            updated.updated_by = request.user
            updated.save(update_fields=[
                'export_planned_date', 'export_date', 'invoice_no',
                'updated_by', 'updated_at',
            ])
    base_url = reverse('expenses:china_export_list')
    if request.POST.get('show') == 'all':
        return redirect(f'{base_url}?show=all')
    return redirect(base_url)


def _serialize_staged_value(value):
    """セッション(JSONシリアライザ)に保存できる形へ変換する。"""
    if isinstance(value, Decimal):
        return {'__decimal__': str(value)}
    # datetime は date のサブクラスなので、必ず date より先に判定する
    if isinstance(value, datetime):
        return {'__datetime__': value.isoformat()}
    if isinstance(value, date):
        return {'__date__': value.isoformat()}
    return value


def _deserialize_staged_value(value):
    if isinstance(value, dict):
        if '__decimal__' in value:
            return Decimal(value['__decimal__'])
        if '__datetime__' in value:
            return datetime.fromisoformat(value['__datetime__'])
        if '__date__' in value:
            return date.fromisoformat(value['__date__'])
    return value


_UPLOAD_TABLE_NAME = 't_china_export'
_UPLOAD_SESSION_KEY = 'china_export_upload_staged'
_UPLOAD_MAX_ROWS = 2000


@login_required
def china_export_upload(request):
    """中国輸出実績報告: Excelファイルのドラッグ&ドロップアップロード（プレビュー段階）。
    ファイルはメモリ上でのみ処理し、ディスクへは保存しない。"""
    _require_china_export_access(request.user)
    # GET/POST を問わず、まず以前ステージしたデータを破棄する。
    # 今回のPOSTの解析・検証が成功した場合にのみ再度ステージされる。
    request.session.pop(_UPLOAD_SESSION_KEY, None)

    if request.method == 'GET':
        return render(request, 'expenses/china_export_upload.html', {'current': 'china_export_list'})

    upload_file = request.FILES.get('excel_file')
    if not upload_file or not upload_file.name.lower().endswith('.xlsx'):
        return render(request, 'expenses/china_export_upload.html', {
            'current': 'china_export_list',
            'file_error': '対応形式は.xlsxのみです。',
        })

    mapping = get_field_mapping(_UPLOAD_TABLE_NAME)
    if not mapping:
        return render(request, 'expenses/china_export_upload.html', {
            'current': 'china_export_list',
            'file_error': '見出し変換マスタが未設定です。管理者に「マスタ設定」からの登録を依頼してください。',
        })

    try:
        valid_rows, errors = parse_excel_rows(upload_file, mapping, T_ChinaExport)
    except Exception:
        logger.exception('中国輸出実績報告アップロード: ファイル読み込みに失敗')
        return render(request, 'expenses/china_export_upload.html', {
            'current': 'china_export_list',
            'file_error': 'ファイルの読み込みに失敗しました。ファイル形式をご確認ください。',
        })

    if errors:
        return render(request, 'expenses/china_export_upload.html', {
            'current': 'china_export_list',
            'errors': errors,
        })

    if len(valid_rows) > _UPLOAD_MAX_ROWS:
        return render(request, 'expenses/china_export_upload.html', {
            'current': 'china_export_list',
            'file_error': f'一度にアップロードできるのは{_UPLOAD_MAX_ROWS}件までです（{len(valid_rows)}件検出）。ファイルを分割してアップロードしてください。',
        })

    request.session[_UPLOAD_SESSION_KEY] = [
        {k: _serialize_staged_value(v) for k, v in row.items()} for row in valid_rows
    ]
    return render(request, 'expenses/china_export_upload.html', {
        'current': 'china_export_list',
        'preview_rows': valid_rows,
        'preview_count': len(valid_rows),
    })


@login_required
@require_POST
def china_export_upload_confirm(request):
    """プレビューで検証済みのデータ(セッション)を確定保存する。"""
    _require_china_export_access(request.user)
    # 二重送信でバッチが重複挿入されるのを防ぐため、読み出しと同時にセッションから除去する
    staged = request.session.pop(_UPLOAD_SESSION_KEY, None)
    if not staged:
        messages.error(request, 'アップロードするデータがありません。ファイルを再度アップロードしてください。')
        return redirect('expenses:china_export_upload')

    records = [
        T_ChinaExport(**{k: _deserialize_staged_value(v) for k, v in row.items()})
        for row in staged
    ]
    with transaction.atomic():
        T_ChinaExport.objects.bulk_create(records, batch_size=500)
    messages.success(request, f'{len(records)}件を取り込みました。')
    return redirect('expenses:china_export_list')
