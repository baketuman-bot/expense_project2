"""中国輸出Invoice管理: ビューのテスト"""
import io
import os
from datetime import date
from decimal import Decimal
from urllib.parse import quote

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from expenses.models import M_Item, M_UserRole, T_ChinaInvoice, T_ChinaInvoiceMonthClose

User = get_user_model()


def _make_users():
    reporter = User.objects.create_user(
        username='view_reporter', man_number='9401', user_name='view報告者', password='pass')
    M_UserRole.objects.create(man_number=reporter, role='china_reporter')
    other = User.objects.create_user(
        username='view_other', man_number='9402', user_name='view権限なし', password='pass')
    accountant = User.objects.create_user(
        username='view_accountant', man_number='9403', user_name='view経理', password='pass')
    M_UserRole.objects.create(man_number=accountant, role='accountant')
    admin = User.objects.create_user(
        username='view_admin', man_number='9404', user_name='view管理者', password='pass')
    M_UserRole.objects.create(man_number=admin, role='admin')
    return reporter, other, accountant, admin


def _make_masters():
    cargo = M_Item.objects.create(data_kbn='CHN_CARGO', key='v1', content='製品', content2='')
    adjrate = M_Item.objects.create(data_kbn='CHN_ADJRT', key='v1', content='0%', content2='0.00')
    return cargo, adjrate


class ChinaInvoiceListViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.cargo, cls.adjrate = _make_masters()
        cls.record1 = T_ChinaInvoice.objects.create(
            invoice_no='INV-LIST-1', invoice_total=Decimal('1000.00'), export_date=date(2026, 8, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i1.pdf', b'a'), reporter=cls.reporter,
        )
        cls.record2 = T_ChinaInvoice.objects.create(
            invoice_no='INV-LIST-2', invoice_total=Decimal('2000.00'), export_date=date(2026, 8, 2),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i2.pdf', b'b'), reporter=cls.reporter,
            accounting_confirmed=True,
        )

    def test_権限がないユーザーは403(self):
        self.client.force_login(self.other)
        res = self.client.get(reverse('expenses:china_invoice_list'))
        self.assertEqual(res.status_code, 403)

    def test_一覧に全件表示される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_list'))
        self.assertContains(res, 'INV-LIST-1')
        self.assertContains(res, 'INV-LIST-2')

    def test_invoice_noで絞り込める(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_list') + '?invoice_no=LIST-1')
        self.assertContains(res, 'INV-LIST-1')
        self.assertNotContains(res, 'INV-LIST-2')

    def test_経理確認状況で絞り込める(self):
        self.client.force_login(self.accountant)
        res = self.client.get(reverse('expenses:china_invoice_list') + '?accounting_confirmed=1')
        self.assertContains(res, 'INV-LIST-2')
        self.assertNotContains(res, 'INV-LIST-1')

    def test_invoice_totalに数値でない文字列を渡しても500にならない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_list') + '?invoice_total=abc')
        self.assertEqual(res.status_code, 200)

    def test_invoice_totalにカンマ区切りの数値で絞り込める(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_list') + '?invoice_total=1,000')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'INV-LIST-1')
        self.assertNotContains(res, 'INV-LIST-2')

    def test_invoice_totalの完全一致で絞り込める(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_list') + '?invoice_total=1000.00')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'INV-LIST-1')
        self.assertNotContains(res, 'INV-LIST-2')


class ChinaInvoiceDetailEditViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.reporter2 = User.objects.create_user(
            username='view_reporter2', man_number='9405', user_name='view報告者2', password='pass')
        M_UserRole.objects.create(man_number=cls.reporter2, role='china_reporter')
        cls.cargo, cls.adjrate = _make_masters()

    def _make_record(self, **overrides):
        data = dict(
            invoice_no='INV-EDIT-1', invoice_total=Decimal('999.00'), export_date=date(2026, 8, 1),
            cargo_category=self.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=self.reporter,
        )
        data.update(overrides)
        return T_ChinaInvoice.objects.create(**data)

    def test_詳細画面が表示される(self):
        record = self._make_record()
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_detail', args=[record.pk]))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'INV-EDIT-1')

    def test_経理確認前は本人の報告者が編集できる(self):
        record = self._make_record()
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'INV-EDIT-2', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[record.pk]))
        record.refresh_from_db()
        self.assertEqual(record.invoice_no, 'INV-EDIT-2')

    def test_他人の報告者は編集できない(self):
        record = self._make_record()
        self.client.force_login(self.reporter2)
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'HACKED', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        self.assertEqual(res.status_code, 403)

    def test_経理確認後は報告者が編集できない(self):
        record = self._make_record(accounting_confirmed=True)
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'HACKED', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        self.assertEqual(res.status_code, 403)

    def test_経理確認後でも経理担当者は編集できる(self):
        record = self._make_record(accounting_confirmed=True)
        self.client.force_login(self.accountant)
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'INV-EDIT-3', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[record.pk]))
        record.refresh_from_db()
        self.assertEqual(record.invoice_no, 'INV-EDIT-3')

    def test_主要項目を変更すると経理確認と中国側確認がリセットされる(self):
        record = self._make_record(
            accounting_confirmed=True, china_confirm_status=T_ChinaInvoice.CHINA_STATUS_CONFIRMED)
        self.client.force_login(self.accountant)
        self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'INV-EDIT-CHANGED', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        record.refresh_from_db()
        self.assertFalse(record.accounting_confirmed)
        self.assertEqual(record.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED)

    def test_ファイルを差し替えずに保存してもファイルは維持される(self):
        record = self._make_record()
        original_name = record.invoice_file.name
        self.client.force_login(self.reporter)
        self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'INV-EDIT-4', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        record.refresh_from_db()
        self.assertEqual(record.invoice_file.name, original_name)

    def test_添付ファイルのみの変更では確認状態が維持される(self):
        record = self._make_record(
            accounting_confirmed=True, china_confirm_status=T_ChinaInvoice.CHINA_STATUS_CONFIRMED)
        self.client.force_login(self.accountant)
        new_file = SimpleUploadedFile('new_invoice.pdf', b'new')
        self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': record.invoice_no, 'invoice_total': str(record.invoice_total),
            'export_date': record.export_date.isoformat(),
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
            'invoice_file': new_file,
        })
        record.refresh_from_db()
        self.assertTrue(record.accounting_confirmed)
        self.assertEqual(record.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_CONFIRMED)

    def test_ファイル差替時に旧ファイルが削除される(self):
        record = self._make_record()
        old_path = record.invoice_file.path
        self.assertTrue(os.path.exists(old_path))
        self.client.force_login(self.reporter)
        new_file = SimpleUploadedFile('replaced.pdf', b'new content')
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': record.invoice_no, 'invoice_total': str(record.invoice_total),
            'export_date': record.export_date.isoformat(),
            'cargo_category': self.cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
            'invoice_file': new_file,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[record.pk]))
        self.assertFalse(os.path.exists(old_path))
        record.refresh_from_db()
        self.assertTrue(os.path.exists(record.invoice_file.path))

    def test_不正な入力後の再表示では読取専用サマリに未保存の値が表示されない(self):
        record = self._make_record()
        self.client.force_login(self.reporter)
        other_cargo = M_Item.objects.create(data_kbn='CHN_CARGO', key='inv1', content='その他', content2='OTHER')
        res = self.client.post(reverse('expenses:china_invoice_detail', args=[record.pk]), {
            'invoice_no': 'SHOULD-NOT-LEAK', 'invoice_total': '999.00', 'export_date': '2026-08-01',
            'cargo_category': other_cargo.pk, 'cargo_note': '', 'adjustment_rate_item': self.adjrate.pk,
        })
        self.assertEqual(res.status_code, 200)
        # 読取専用サマリ表示用のinvoiceはDBの値のまま（未保存の入力でin-place変更されていない）
        self.assertEqual(res.context['invoice'].invoice_no, 'INV-EDIT-1')
        # 編集フォーム側は差戻し用に入力値をそのまま保持している（通常の挙動）
        self.assertEqual(res.context['form'].data['invoice_no'], 'SHOULD-NOT-LEAK')
        record.refresh_from_db()
        self.assertEqual(record.invoice_no, 'INV-EDIT-1')


class ChinaInvoicePackingListViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.cargo, cls.adjrate = _make_masters()
        cls.record = T_ChinaInvoice.objects.create(
            invoice_no='INV-PL-1', invoice_total=Decimal('1.00'), export_date=date(2026, 8, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=cls.reporter,
        )

    def test_報告者はPacking_Listを追加できる(self):
        self.client.force_login(self.reporter)
        res = self.client.post(
            reverse('expenses:china_invoice_packing_list_add', args=[self.record.pk]),
            {'packing_list_files': SimpleUploadedFile('pl.pdf', b'x')})
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[self.record.pk]))
        self.assertEqual(self.record.packing_lists.count(), 1)

    def test_報告者はPacking_Listを削除できる(self):
        from expenses.models import T_ChinaInvoicePackingList
        pl = T_ChinaInvoicePackingList.objects.create(
            invoice=self.record, file=SimpleUploadedFile('pl.pdf', b'x'), uploaded_by=self.reporter)
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_packing_list_delete', args=[pl.pk]))
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[self.record.pk]))
        self.assertEqual(self.record.packing_lists.count(), 0)

    def test_不正なファイルはPacking_List追加で500にならず作成されない(self):
        self.client.force_login(self.reporter)
        res = self.client.post(
            reverse('expenses:china_invoice_packing_list_add', args=[self.record.pk]),
            {'packing_list_files': SimpleUploadedFile('bad.txt', b'x')})
        self.assertRedirects(res, reverse('expenses:china_invoice_detail', args=[self.record.pk]))
        self.assertEqual(self.record.packing_lists.count(), 0)


class ChinaInvoiceDeleteViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.reporter2 = User.objects.create_user(
            username='view_reporter3', man_number='9406', user_name='view報告者3', password='pass')
        M_UserRole.objects.create(man_number=cls.reporter2, role='china_reporter')
        cls.cargo, cls.adjrate = _make_masters()

    def _make_record(self, **overrides):
        data = dict(
            invoice_no='INV-DEL-1', invoice_total=Decimal('1.00'), export_date=date(2026, 8, 1),
            cargo_category=self.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=self.reporter,
        )
        data.update(overrides)
        return T_ChinaInvoice.objects.create(**data)

    def test_経理確認前は本人の報告者が削除できる(self):
        record = self._make_record()
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertRedirects(res, reverse('expenses:china_invoice_list'))
        self.assertFalse(T_ChinaInvoice.objects.filter(pk=record.pk).exists())

    def test_他人の報告者は削除できない(self):
        record = self._make_record()
        self.client.force_login(self.reporter2)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertEqual(res.status_code, 403)
        self.assertTrue(T_ChinaInvoice.objects.filter(pk=record.pk).exists())

    def test_経理確認後は報告者が削除できない(self):
        record = self._make_record(accounting_confirmed=True)
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertEqual(res.status_code, 403)

    def test_経理確認後でも経理担当者は削除できる(self):
        record = self._make_record(accounting_confirmed=True)
        self.client.force_login(self.accountant)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertRedirects(res, reverse('expenses:china_invoice_list'))
        self.assertFalse(T_ChinaInvoice.objects.filter(pk=record.pk).exists())

    def test_china_partnerは削除できない(self):
        partner = User.objects.create_user(
            username='view_partner1', man_number='9407', user_name='view中国側1', password='pass')
        M_UserRole.objects.create(man_number=partner, role='china_partner')
        record = self._make_record()
        self.client.force_login(partner)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertEqual(res.status_code, 403)

    def test_削除時にPacking_Listのファイルも削除される(self):
        from expenses.models import T_ChinaInvoicePackingList
        record = self._make_record()
        pl = T_ChinaInvoicePackingList.objects.create(
            invoice=record, file=SimpleUploadedFile('pl.pdf', b'x'), uploaded_by=self.reporter)
        pl_path = pl.file.path
        self.assertTrue(os.path.exists(pl_path))
        self.client.force_login(self.reporter)
        res = self.client.post(reverse('expenses:china_invoice_delete', args=[record.pk]))
        self.assertRedirects(res, reverse('expenses:china_invoice_list'))
        self.assertFalse(os.path.exists(pl_path))


class ChinaInvoiceAccountingViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.cargo, cls.adjrate = _make_masters()
        cls.unconfirmed1 = T_ChinaInvoice.objects.create(
            invoice_no='INV-ACC-1', invoice_total=Decimal('1.00'), export_date=date(2026, 8, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=cls.reporter,
        )
        cls.unconfirmed2 = T_ChinaInvoice.objects.create(
            invoice_no='INV-ACC-2', invoice_total=Decimal('2.00'), export_date=date(2026, 8, 2),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i2.pdf', b'b'), reporter=cls.reporter,
        )

    def test_reporterロールだけでは経理確認画面にアクセスできない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_accounting'))
        self.assertEqual(res.status_code, 403)

    def test_accountantは経理確認画面にアクセスできる(self):
        self.client.force_login(self.accountant)
        res = self.client.get(reverse('expenses:china_invoice_accounting'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'INV-ACC-1')

    def test_個別に確認済みにできる(self):
        self.client.force_login(self.accountant)
        res = self.client.post(reverse('expenses:china_invoice_accounting_confirm'), {
            'pks': [self.unconfirmed1.pk],
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_accounting'))
        self.unconfirmed1.refresh_from_db()
        self.assertTrue(self.unconfirmed1.accounting_confirmed)
        self.assertEqual(self.unconfirmed1.accounting_confirmed_by, self.accountant)
        self.unconfirmed2.refresh_from_db()
        self.assertFalse(self.unconfirmed2.accounting_confirmed)

    def test_複数選択で一括確認できる(self):
        self.client.force_login(self.accountant)
        self.client.post(reverse('expenses:china_invoice_accounting_confirm'), {
            'pks': [self.unconfirmed1.pk, self.unconfirmed2.pk],
        })
        self.unconfirmed1.refresh_from_db()
        self.unconfirmed2.refresh_from_db()
        self.assertTrue(self.unconfirmed1.accounting_confirmed)
        self.assertTrue(self.unconfirmed2.accounting_confirmed)

    def test_未確認をすべて確認済みにできる(self):
        self.client.force_login(self.accountant)
        self.client.post(reverse('expenses:china_invoice_accounting_confirm'), {'confirm_all': '1'})
        self.unconfirmed1.refresh_from_db()
        self.unconfirmed2.refresh_from_db()
        self.assertTrue(self.unconfirmed1.accounting_confirmed)
        self.assertTrue(self.unconfirmed2.accounting_confirmed)

    def test_経理確認操作は中国側確認状態に影響しない(self):
        self.unconfirmed1.china_confirm_status = T_ChinaInvoice.CHINA_STATUS_DIFFERENCE
        self.unconfirmed1.save(update_fields=['china_confirm_status'])
        self.client.force_login(self.accountant)
        self.client.post(reverse('expenses:china_invoice_accounting_confirm'), {
            'pks': [self.unconfirmed1.pk, self.unconfirmed2.pk],
        })
        self.unconfirmed1.refresh_from_db()
        self.unconfirmed2.refresh_from_db()
        self.assertTrue(self.unconfirmed1.accounting_confirmed)
        self.assertEqual(self.unconfirmed1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_DIFFERENCE)
        self.assertEqual(self.unconfirmed2.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED)

    def test_china_partnerロールでは経理確認画面にアクセスできない(self):
        partner = User.objects.create_user(
            username='view_partner_acc', man_number='9410', user_name='view中国側acc', password='pass')
        M_UserRole.objects.create(man_number=partner, role='china_partner')
        self.client.force_login(partner)
        res = self.client.get(reverse('expenses:china_invoice_accounting'))
        self.assertEqual(res.status_code, 403)


class ChinaInvoiceMonthCloseViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()

    def test_reporterロールだけでは月締め画面にアクセスできない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_month_close'))
        self.assertEqual(res.status_code, 403)

    def test_accountantは月を締められる(self):
        self.client.force_login(self.accountant)
        res = self.client.post(reverse('expenses:china_invoice_month_close'), {'year_month': '2026-08'})
        self.assertRedirects(res, reverse('expenses:china_invoice_month_close'))
        self.assertTrue(T_ChinaInvoiceMonthClose.objects.filter(year_month='2026-08').exists())

    def test_既に締めた月は再度締められない(self):
        T_ChinaInvoiceMonthClose.objects.create(year_month='2026-08', closed_by=self.accountant)
        self.client.force_login(self.accountant)
        res = self.client.post(reverse('expenses:china_invoice_month_close'), {'year_month': '2026-08'}, follow=True)
        self.assertContains(res, '既に締め済み')
        self.assertEqual(T_ChinaInvoiceMonthClose.objects.filter(year_month='2026-08').count(), 1)


class ChinaInvoiceChinaCheckViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.partner = User.objects.create_user(
            username='view_partner2', man_number='9408', user_name='view中国側2', password='pass')
        M_UserRole.objects.create(man_number=cls.partner, role='china_partner')
        cls.cargo, cls.adjrate = _make_masters()
        cls.record1 = T_ChinaInvoice.objects.create(
            invoice_no='INV-CC-1', invoice_total=Decimal('1.00'), export_date=date(2026, 8, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=cls.reporter,
        )
        cls.record2 = T_ChinaInvoice.objects.create(
            invoice_no='INV-CC-2', invoice_total=Decimal('2.00'), export_date=date(2026, 8, 2),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i2.pdf', b'b'), reporter=cls.reporter,
        )

    def test_china_partner以外はアクセスできない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_china_check'))
        self.assertEqual(res.status_code, 403)

    def test_china_partnerは一覧を閲覧できる(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_china_check'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'INV-CC-1')

    def test_個別に差異ありへ変更できる(self):
        self.client.force_login(self.partner)
        res = self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pk': self.record1.pk, 'status': T_ChinaInvoice.CHINA_STATUS_DIFFERENCE,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_china_check'))
        self.record1.refresh_from_db()
        self.assertEqual(self.record1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_DIFFERENCE)
        self.assertEqual(self.record1.china_confirmed_by, self.partner)

    def test_一括確認済みにできる(self):
        self.client.force_login(self.partner)
        self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pks': [self.record1.pk, self.record2.pk], 'bulk_status': T_ChinaInvoice.CHINA_STATUS_CONFIRMED,
        })
        self.record1.refresh_from_db()
        self.record2.refresh_from_db()
        self.assertEqual(self.record1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_CONFIRMED)
        self.assertEqual(self.record2.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_CONFIRMED)

    def test_一括で差異ありには変更できない(self):
        self.client.force_login(self.partner)
        res = self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pks': [self.record1.pk], 'bulk_status': T_ChinaInvoice.CHINA_STATUS_DIFFERENCE,
        })
        self.assertEqual(res.status_code, 400)
        self.record1.refresh_from_db()
        self.assertEqual(self.record1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED)

    def test_日本側の経理確認状態は中国側の操作で変わらない(self):
        self.record1.accounting_confirmed = True
        self.record1.save(update_fields=['accounting_confirmed'])
        self.client.force_login(self.partner)
        self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pk': self.record1.pk, 'status': T_ChinaInvoice.CHINA_STATUS_DIFFERENCE,
        })
        self.record1.refresh_from_db()
        self.assertTrue(self.record1.accounting_confirmed)

    def test_個別に確認済みへ変更できる(self):
        self.client.force_login(self.partner)
        res = self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pk': self.record1.pk, 'status': T_ChinaInvoice.CHINA_STATUS_CONFIRMED,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_china_check'))
        self.record1.refresh_from_db()
        self.assertEqual(self.record1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_CONFIRMED)
        self.assertEqual(self.record1.china_confirmed_by, self.partner)

    def test_個別に未確認へ戻せる(self):
        self.record1.china_confirm_status = T_ChinaInvoice.CHINA_STATUS_CONFIRMED
        self.record1.save(update_fields=['china_confirm_status'])
        self.client.force_login(self.partner)
        res = self.client.post(reverse('expenses:china_invoice_china_check_update'), {
            'pk': self.record1.pk, 'status': T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED,
        })
        self.assertRedirects(res, reverse('expenses:china_invoice_china_check'))
        self.record1.refresh_from_db()
        self.assertEqual(self.record1.china_confirm_status, T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED)
        self.assertEqual(self.record1.china_confirmed_by, self.partner)

    def test_一覧に貨物概要補足など詳細項目は表示されない(self):
        self.record1.cargo_note = 'SENTINEL_NOTE_VALUE'
        self.record1.save(update_fields=['cargo_note'])
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_china_check'))
        self.assertNotContains(res, 'SENTINEL_NOTE_VALUE')

    def test_登録日で絞り込める(self):
        T_ChinaInvoice.objects.filter(pk=self.record1.pk).update(registered_at='2026-08-05 09:00:00')
        T_ChinaInvoice.objects.filter(pk=self.record2.pk).update(registered_at='2026-08-06 09:00:00')
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_china_check') + '?registered_date=2026-08-05')
        self.assertContains(res, 'INV-CC-1')
        self.assertNotContains(res, 'INV-CC-2')

    def test_輸出日で絞り込める(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_china_check') + '?export_date=2026-08-01')
        self.assertContains(res, 'INV-CC-1')
        self.assertNotContains(res, 'INV-CC-2')

    def test_対象年月で絞り込める(self):
        T_ChinaInvoice.objects.filter(pk=self.record1.pk).update(registered_at='2026-08-05 09:00:00')
        T_ChinaInvoice.objects.filter(pk=self.record2.pk).update(registered_at='2026-09-15 09:00:00')
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_china_check') + '?month=2026-09')
        self.assertContains(res, 'INV-CC-2')
        self.assertNotContains(res, 'INV-CC-1')


class ChinaInvoiceExcelViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.cargo, cls.adjrate = _make_masters()
        cls.record_b = T_ChinaInvoice.objects.create(
            invoice_no='INV-XL-B', invoice_total=Decimal('200.00'), export_date=date(2026, 8, 2),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=cls.reporter,
        )
        cls.record_a = T_ChinaInvoice.objects.create(
            invoice_no='INV-XL-A', invoice_total=Decimal('100.00'), export_date=date(2026, 8, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i2.pdf', b'b'), reporter=cls.reporter,
        )
        cls.record_c = T_ChinaInvoice.objects.create(
            invoice_no='INV-XL-C', invoice_total=Decimal('300.00'), export_date=date(2026, 9, 1),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i3.pdf', b'c'), reporter=cls.reporter,
        )
        # registered_at は auto_now_add のため .create() では上書きできない。
        # .update() は auto_now_add の save()-time 上書きをバイパスするため、
        # year_month/date_from-date_to フィルタ用に意図的に2026-08範囲外の値へ変更する。
        T_ChinaInvoice.objects.filter(pk=cls.record_c.pk).update(registered_at='2026-09-15 10:00:00')
        cls.record_c.refresh_from_db()

    def test_権限がなければ403(self):
        self.client.force_login(self.other)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        self.assertEqual(res.status_code, 403)

    # 帳票レイアウト: 1-2行目=タイトル・検印欄、4行目=表ヘッダー、5行目〜=データ、最終行=合計
    HEADER_ROW = 4
    DATA_START_ROW = 5

    def test_Invoice_No昇順で出力される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        invoice_no_col_values = [
            row[1].value for row in ws.iter_rows(min_row=self.DATA_START_ROW, max_row=ws.max_row - 1)]
        self.assertEqual(invoice_no_col_values, ['INV-XL-A', 'INV-XL-B', 'INV-XL-C'])

    def test_月単位のファイル名になる(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel') + '?year_month=2026-08')
        # 日本語ファイル名はRFC 5987 (filename*=UTF-8''...) でパーセントエンコードされて出力される
        # （expenses/views.py のCSV出力と同じ content_disposition_header() を使うため）
        self.assertIn(quote('中国輸出実績_202608.xlsx'), res['Content-Disposition'])

    def test_年月指定で範囲外レコードは除外される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel') + '?year_month=2026-08')
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        invoice_no_col_values = [row[1].value for row in ws.iter_rows(min_row=self.DATA_START_ROW)]
        self.assertIn('INV-XL-A', invoice_no_col_values)
        self.assertIn('INV-XL-B', invoice_no_col_values)
        self.assertNotIn('INV-XL-C', invoice_no_col_values)

    def test_日付範囲指定で絞り込める(self):
        self.client.force_login(self.reporter)
        res = self.client.get(
            reverse('expenses:china_invoice_excel') + '?date_from=2026-09-01&date_to=2026-09-30')
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        invoice_no_col_values = [
            row[1].value for row in ws.iter_rows(min_row=self.DATA_START_ROW, max_row=ws.max_row - 1)]
        self.assertEqual(invoice_no_col_values, ['INV-XL-C'])
        self.assertIn(quote('中国輸出実績_20260901-20260930.xlsx'), res['Content-Disposition'])

    def test_同じ月は同じファイル名になる(self):
        self.client.force_login(self.reporter)
        res1 = self.client.get(reverse('expenses:china_invoice_excel') + '?year_month=2026-08')
        res2 = self.client.get(reverse('expenses:china_invoice_excel') + '?year_month=2026-08')
        self.assertEqual(res1['Content-Disposition'], res2['Content-Disposition'])

    def test_Invoiceファイル等の列は含まれない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        header = [c.value for c in wb.active[self.HEADER_ROW]]
        self.assertNotIn('Invoiceファイル', header)
        self.assertNotIn('Packing List', header)
        self.assertNotIn('中国側確認', header)
        self.assertNotIn('経理確認', header)

    def test_ヘッダーは経理確認と中国確認を除く全項目(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        header = [c.value for c in wb.active[self.HEADER_ROW]]
        self.assertEqual(header, [
            '管理番号', 'Invoice No', 'Invoice Total', '輸出日', '貨物\n概要\n区分',
            '貨物\n概要\n補足', '加算調整率', '報告者', '登録日時',
            '通貨', '元値相当（通貨）', '管理費（通貨）', '元値相当（JPY）', '管理費（JPY）', '金額（JPY）',
        ])
        # 折り返しヘッダー（貨物概要区分・補足）が表示されるよう wrap_text を有効にしている
        self.assertTrue(wb.active.cell(row=self.HEADER_ROW, column=5).alignment.wrap_text)

    def test_検印欄がE列からG列にある(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # 検印欄: 1行目のE/F/G列に承認・確認・担当のラベル、2行目が押印用の空欄
        self.assertEqual([ws['E1'].value, ws['F1'].value, ws['G1'].value], ['承認', '確認', '担当'])
        self.assertEqual([ws['E2'].value, ws['F2'].value, ws['G2'].value], [None, None, None])
        self.assertIsNotNone(ws['E2'].border.bottom.style)   # 押印枠に罫線がある

    def test_印刷の左右余白が05インチ(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.page_margins.left, 0.5)
        self.assertEqual(ws.page_margins.right, 0.5)

    def test_JPY換算列は為替レートセル参照の数式で出力される(self):
        # INVOICE実績報告書のK〜P列相当。為替レートはN3へ手入力する運用のため数式で出力する
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws['M3'].value, '為替レート')          # レート入力欄のラベル
        self.assertIsNotNone(ws['N3'].border.bottom.style)     # レート入力欄に枠がある
        row = self.DATA_START_ROW
        self.assertEqual(ws[f'J{row}'].value, 'US$')
        self.assertEqual(ws[f'K{row}'].value, f'=C{row}/(1+G{row}/100)')
        self.assertEqual(ws[f'L{row}'].value, f'=C{row}-K{row}')
        self.assertEqual(ws[f'M{row}'].value, f'=IF($N$3="","",O{row}-N{row})')
        self.assertEqual(
            ws[f'N{row}'].value, f'=IF($N$3="","",ROUND(IF(J{row}<>"JPY",$N$3,1)*L{row},0))')
        self.assertEqual(
            ws[f'O{row}'].value, f'=IF($N$3="","",ROUND(IF(J{row}<>"JPY",$N$3,1)*C{row},0))')
        # 合計行には換算系列のSUM数式が入る
        total_row = ws.max_row
        self.assertEqual(
            ws.cell(row=total_row, column=11).value,
            f'=SUM(K{self.DATA_START_ROW}:K{total_row - 1})')

    def test_タイトルと対象期間が出力される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel') + '?year_month=2026-08')
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws['A1'].value, '中国輸出実績報告')
        self.assertIn('2026-08', ws['A2'].value)

    def test_最終行に件数付きの合計行が出力される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        total_row = ws.max_row
        self.assertEqual(ws.cell(row=total_row, column=1).value, '合計（3件）')
        # Invoice Total (C列) の合計: 100 + 200 + 300
        self.assertEqual(ws.cell(row=total_row, column=3).value, 600.0)
        self.assertTrue(ws.cell(row=total_row, column=1).font.bold)

    def test_表ヘッダーに色があり偶数データ行が縞模様になる(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=self.HEADER_ROW, column=1).fill.fill_type, 'solid')
        # 1行目のデータ行は無地、2行目のデータ行に縞色が付く
        self.assertIsNone(ws.cell(row=self.DATA_START_ROW, column=1).fill.fill_type)
        self.assertEqual(ws.cell(row=self.DATA_START_ROW + 1, column=1).fill.fill_type, 'solid')

    def test_データ表には罫線がない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        header_cell = ws.cell(row=self.HEADER_ROW, column=1)
        data_cell = ws.cell(row=self.DATA_START_ROW, column=1)
        for cell in (header_cell, data_cell):
            self.assertIsNone(cell.border.top.style)
            self.assertIsNone(cell.border.bottom.style)
            self.assertIsNone(cell.border.left.style)
            self.assertIsNone(cell.border.right.style)
        self.assertFalse(ws.sheet_view.showGridLines)


class ChinaInvoiceDashboardViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.cargo, cls.adjrate = _make_masters()
        T_ChinaInvoice.objects.create(
            invoice_no='INV-DASH-1', invoice_total=Decimal('1.00'), export_date=date.today(),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i.pdf', b'a'), reporter=cls.reporter,
            china_confirm_status=T_ChinaInvoice.CHINA_STATUS_DIFFERENCE,
        )
        cls.confirmed_other_record = T_ChinaInvoice.objects.create(
            invoice_no='INV-DASH-2', invoice_total=Decimal('2.00'), export_date=date.today(),
            cargo_category=cls.cargo, adjustment_rate_value=Decimal('0.00'),
            invoice_file=SimpleUploadedFile('i2.pdf', b'b'), reporter=cls.reporter,
            accounting_confirmed=True, china_confirm_status=T_ChinaInvoice.CHINA_STATUS_CONFIRMED,
        )

    def test_権限がなければ403(self):
        self.client.force_login(self.other)
        res = self.client.get(reverse('expenses:china_invoice_dashboard'))
        self.assertEqual(res.status_code, 403)

    def test_サマリ件数が表示される(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['unconfirmed_accounting_count'], 1)
        self.assertEqual(res.context['difference_count'], 1)
        self.assertEqual(res.context['this_month_count'], 2)


class ChinaInvoiceRoleGatedButtonTests(TestCase):
    """ダッシュボード・一覧のInvoice登録ボタン/経理確認リンクが権限に応じて非表示になることの確認"""

    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.partner = User.objects.create_user(
            username='view_partner_gate', man_number='9411', user_name='viewゲート', password='pass')
        M_UserRole.objects.create(man_number=cls.partner, role='china_partner')

    def test_china_partnerのみのユーザーにはダッシュボードでInvoice報告リンクが出ない(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_dashboard'))
        self.assertNotContains(res, reverse('expenses:china_invoice_report_upload'))

    def test_china_partnerのみのユーザーには一覧でInvoice報告リンクが出ない(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_list'))
        self.assertNotContains(res, reverse('expenses:china_invoice_report_upload'))

    def test_china_partnerのみのユーザーにはダッシュボードで経理確認へのリンクが出ない(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:china_invoice_dashboard'))
        self.assertNotContains(res, reverse('expenses:china_invoice_accounting'))


class ChinaInvoiceSidebarTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.reporter, cls.other, cls.accountant, cls.admin = _make_users()
        cls.partner = User.objects.create_user(
            username='view_partner3', man_number='9409', user_name='view中国側3', password='pass')
        M_UserRole.objects.create(man_number=cls.partner, role='china_partner')

    def test_china_reporterはサイドバーにメニューが出る(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, '中国輸出Invoice管理')

    def test_china_partnerはサイドバーにメニューが出る(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, '中国輸出Invoice管理')

    def test_権限がないユーザーには出ない(self):
        self.client.force_login(self.other)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, '中国輸出Invoice管理')

    def test_reporterのメニューにはInvoice報告リンクがある(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, reverse('expenses:china_invoice_report_upload'))

    def test_中国側ユーザーには経理確認リンクは出ない(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, reverse('expenses:china_invoice_accounting'))

    def test_accountantはサイドバーにメニューが出る(self):
        self.client.force_login(self.accountant)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, '中国輸出Invoice管理')

    def test_reporterには経理確認_月締め_中国側確認リンクが出ない(self):
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, reverse('expenses:china_invoice_accounting'))
        self.assertNotContains(res, reverse('expenses:china_invoice_month_close'))
        self.assertNotContains(res, reverse('expenses:china_invoice_china_check'))

    def test_accountantにはInvoice報告_中国側確認リンクが出ない(self):
        self.client.force_login(self.accountant)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, reverse('expenses:china_invoice_report_upload'))
        self.assertNotContains(res, reverse('expenses:china_invoice_china_check'))

    def test_partnerにはInvoice報告_経理確認_月締めリンクが出ない(self):
        self.client.force_login(self.partner)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, reverse('expenses:china_invoice_report_upload'))
        self.assertNotContains(res, reverse('expenses:china_invoice_accounting'))
        self.assertNotContains(res, reverse('expenses:china_invoice_month_close'))


class ChinaInvoiceFileOptionalTests(TestCase):
    """invoice_file 任意化（パッキングリストExcel取込対応）"""

    @classmethod
    def setUpTestData(cls):
        cls.reporter = User.objects.create_user(
            username='opt_reporter', man_number='9701', user_name='opt報告者', password='pass')
        M_UserRole.objects.create(man_number=cls.reporter, role='china_reporter')
        cls.cargo = M_Item.objects.create(
            data_kbn='CHN_CARGO', key='o1', content='製品', content2='')

    def test_invoice_fileなしで保存できる(self):
        invoice = T_ChinaInvoice.objects.create(
            invoice_no='NOFILE-1', invoice_total=Decimal('10.00'),
            export_date=date(2026, 7, 1), cargo_category=self.cargo,
            adjustment_rate_value=Decimal('0.00'), reporter=self.reporter,
        )
        invoice.refresh_from_db()
        self.assertEqual(invoice.invoice_file.name, '')

    def test_ファイルなしInvoiceの詳細画面が開けてダウンロードリンクが出ない(self):
        invoice = T_ChinaInvoice.objects.create(
            invoice_no='NOFILE-2', invoice_total=Decimal('10.00'),
            export_date=date(2026, 7, 1), cargo_category=self.cargo,
            adjustment_rate_value=Decimal('0.00'), reporter=self.reporter,
        )
        self.client.force_login(self.reporter)
        res = self.client.get(reverse('expenses:china_invoice_detail', args=[invoice.pk]))
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, 'ダウンロード')
        self.assertContains(res, 'なし（Excel取込）')
