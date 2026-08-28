"""中国輸出Invoice管理: Invoice登録・一覧・確認・月締め・Excel出力のビュー。
既存の中国輸出実績報告(T_ChinaExport)とは独立したサブシステム。"""
import datetime
import logging
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_POST

from .china_invoice_files import validate_china_invoice_file
from .forms import ChinaInvoiceForm
from .models import M_Item, M_User, T_ChinaInvoice, T_ChinaInvoiceMonthClose, T_ChinaInvoicePackingList

logger = logging.getLogger(__name__)


def _require_role(user, *roles):
    if user.has_role('admin'):
        return
    if any(user.has_role(r) for r in roles):
        return
    raise PermissionDenied()


def _is_month_closed(target_date):
    return T_ChinaInvoiceMonthClose.objects.filter(year_month=target_date.strftime('%Y-%m')).exists()


def _validate_packing_list_uploads(request, field_name='packing_list_files', with_filename=False):
    """アップロードされた全Packing Listファイルを検証する。
    エラーメッセージのリストを返す（空リスト=全ファイル有効）。
    with_filename=True のとき 'ファイル名: ' の接頭辞を付ける（報告ウィザードのように
    1画面で複数行分を扱う場合に、どの行のファイルかを判別するため）。
    1件でも不正なファイルがあれば呼び出し側で保存処理そのものを中止すること。"""
    errors = []
    for f in request.FILES.getlist(field_name):
        try:
            validate_china_invoice_file(f)
        except ValidationError as e:
            if with_filename:
                errors.extend(f'{f.name}: {m}' for m in e.messages)
            else:
                errors.extend(e.messages)
    if errors:
        logger.warning('Packing Listアップロード検証エラー: %s', errors)
    return errors


def _handle_packing_list_uploads(request, invoice, field_name='packing_list_files'):
    """Packing Listファイルを保存する。呼び出し側で事前に
    _validate_packing_list_uploads() による検証を済ませておくこと（全件有効を前提とする）。"""
    for f in request.FILES.getlist(field_name):
        T_ChinaInvoicePackingList.objects.create(invoice=invoice, file=f, uploaded_by=request.user)


_LIST_ROLES = ('china_reporter', 'accountant', 'china_partner')


def _china_invoice_queryset(request):
    qs = T_ChinaInvoice.objects.select_related('cargo_category', 'reporter').order_by('-registered_at')
    params = request.GET
    if params.get('management_no'):
        qs = qs.filter(management_no__icontains=params['management_no'])
    if params.get('invoice_no'):
        qs = qs.filter(invoice_no__icontains=params['invoice_no'])
    if params.get('export_date'):
        qs = qs.filter(export_date=params['export_date'])
    if params.get('registered_date'):
        qs = qs.filter(registered_at__date=params['registered_date'])
    if params.get('cargo_category'):
        qs = qs.filter(cargo_category_id=params['cargo_category'])
    if params.get('reporter'):
        qs = qs.filter(reporter_id=params['reporter'])
    if params.get('invoice_total'):
        try:
            qs = qs.filter(invoice_total=Decimal(params['invoice_total'].replace(',', '')))
        except InvalidOperation:
            pass
    if params.get('accounting_confirmed') in ('0', '1'):
        qs = qs.filter(accounting_confirmed=(params['accounting_confirmed'] == '1'))
    if params.get('china_confirm_status'):
        qs = qs.filter(china_confirm_status=params['china_confirm_status'])
    if params.get('month_status') in ('closed', 'open'):
        closed_months = set(T_ChinaInvoiceMonthClose.objects.values_list('year_month', flat=True))
        ids = [
            r.pk for r in qs
            if (r.registered_at.strftime('%Y-%m') in closed_months) == (params['month_status'] == 'closed')
        ]
        qs = qs.filter(pk__in=ids)
    return qs


@login_required
def china_invoice_list(request):
    _require_role(request.user, *_LIST_ROLES)
    return render(request, 'expenses/china_invoice_list.html', {
        'records': _china_invoice_queryset(request),
        'cargo_categories': M_Item.objects.filter(data_kbn='CHN_CARGO').order_by('order_by', 'key'),
        'reporters': M_User.objects.filter(roles__role='china_reporter').distinct().order_by('user_name'),
        'china_status_choices': T_ChinaInvoice.CHINA_STATUS_CHOICES,
        'current': 'china_invoice_list',
    })


_KEY_FIELDS = ['invoice_no', 'invoice_total', 'export_date', 'cargo_category_id', 'adjustment_rate_value']


def _can_edit(user, invoice):
    if user.has_role('admin') or user.has_role('accountant'):
        return True
    if user.has_role('china_reporter') and invoice.reporter_id == user.pk and not invoice.accounting_confirmed:
        return True
    return False


def _reset_confirmations_if_key_changed(old_snapshot, new_instance):
    changed = any(old_snapshot[f] != getattr(new_instance, f) for f in _KEY_FIELDS)
    if changed:
        new_instance.accounting_confirmed = False
        new_instance.accounting_confirmed_by = None
        new_instance.accounting_confirmed_at = None
        new_instance.china_confirm_status = T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED
        new_instance.china_confirmed_by = None
        new_instance.china_confirmed_at = None
    return changed


@login_required
def china_invoice_detail(request, pk):
    _require_role(request.user, *_LIST_ROLES)
    invoice = get_object_or_404(
        T_ChinaInvoice.objects.select_related('cargo_category', 'reporter'), pk=pk)
    can_edit = _can_edit(request.user, invoice)
    can_delete = _can_delete(request.user, invoice)
    # 読み取り専用の概要表示に使うインスタンス。POSTが不正だった場合、form.is_valid()が
    # 呼び出し済みのinvoiceを未保存のままin-place変更してしまうため、その値を表示に使わない。
    display_invoice = invoice

    if request.method == 'POST':
        if not can_edit:
            raise PermissionDenied()
        old_snapshot = {f: getattr(invoice, f) for f in _KEY_FIELDS}
        # 差替時に旧ファイルを削除するため、フォームによるinstance書き換え前に参照を保持する
        old_invoice_file = invoice.invoice_file
        form = ChinaInvoiceForm(request.POST, request.FILES, instance=invoice)
        pl_errors = []
        if form.is_valid():
            pl_errors = _validate_packing_list_uploads(request)
            for msg in pl_errors:
                messages.error(request, msg)

        if form.is_valid() and not pl_errors:
            updated = form.save(commit=False)
            if not request.FILES.get('invoice_file'):
                updated.invoice_file = invoice.invoice_file
            elif old_invoice_file:
                # FieldFile.delete()はinstanceのフィールドをNoneに書き換える副作用があり、
                # 直前にconstruct_instance()で設定済みの新ファイルを消してしまうため、
                # instanceに触れないstorage.delete()で旧ファイルのみを削除する。
                old_invoice_file.storage.delete(old_invoice_file.name)
            _reset_confirmations_if_key_changed(old_snapshot, updated)
            updated.save()
            _handle_packing_list_uploads(request, updated)
            messages.success(request, f'{updated.management_no} を更新しました。')
            return redirect('expenses:china_invoice_detail', pk=updated.pk)
        else:
            # 保存されなかった不正データが概要表示に混ざらないよう、DBからクリーンな状態を取り直す
            display_invoice = T_ChinaInvoice.objects.select_related('cargo_category', 'reporter').get(pk=invoice.pk)
    else:
        form = ChinaInvoiceForm(instance=invoice) if can_edit else None

    return render(request, 'expenses/china_invoice_detail.html', {
        'invoice': display_invoice, 'form': form, 'can_edit': can_edit,
        'can_delete': can_delete, 'current': 'china_invoice_list',
    })


@login_required
@require_POST
def china_invoice_packing_list_add(request, pk):
    invoice = get_object_or_404(T_ChinaInvoice, pk=pk)
    if not _can_edit(request.user, invoice):
        raise PermissionDenied()
    pl_errors = _validate_packing_list_uploads(request)
    if pl_errors:
        for msg in pl_errors:
            messages.error(request, msg)
    else:
        _handle_packing_list_uploads(request, invoice)
    return redirect('expenses:china_invoice_detail', pk=invoice.pk)


@login_required
@require_POST
def china_invoice_packing_list_delete(request, pk):
    packing_list = get_object_or_404(T_ChinaInvoicePackingList, pk=pk)
    if not _can_edit(request.user, packing_list.invoice):
        raise PermissionDenied()
    invoice_pk = packing_list.invoice_id
    packing_list.file.delete(save=False)
    packing_list.delete()
    return redirect('expenses:china_invoice_detail', pk=invoice_pk)


def _can_delete(user, invoice):
    if user.has_role('admin') or user.has_role('accountant'):
        return True
    if user.has_role('china_reporter') and invoice.reporter_id == user.pk and not invoice.accounting_confirmed:
        return True
    return False


@login_required
@require_POST
def china_invoice_delete(request, pk):
    invoice = get_object_or_404(T_ChinaInvoice, pk=pk)
    if not _can_delete(request.user, invoice):
        raise PermissionDenied()
    management_no = invoice.management_no
    for pl in invoice.packing_lists.all():
        pl.file.delete(save=False)
    invoice.invoice_file.delete(save=False)
    invoice.delete()
    messages.success(request, f'{management_no} を削除しました。')
    return redirect('expenses:china_invoice_list')


@login_required
def china_invoice_accounting(request):
    _require_role(request.user, 'accountant')
    records = T_ChinaInvoice.objects.filter(accounting_confirmed=False).order_by('registered_at')
    return render(request, 'expenses/china_invoice_accounting.html', {
        'records': records, 'current': 'china_invoice_accounting',
    })


@login_required
@require_POST
def china_invoice_accounting_confirm(request):
    _require_role(request.user, 'accountant')
    if request.POST.get('confirm_all') == '1':
        targets = T_ChinaInvoice.objects.filter(accounting_confirmed=False)
    else:
        pks = request.POST.getlist('pks')
        targets = T_ChinaInvoice.objects.filter(pk__in=pks)
    now = timezone.now()
    updated = targets.update(
        accounting_confirmed=True, accounting_confirmed_by=request.user, accounting_confirmed_at=now)
    messages.success(request, f'{updated}件を経理確認済みにしました。')
    return redirect('expenses:china_invoice_accounting')


@login_required
def china_invoice_month_close(request):
    _require_role(request.user, 'accountant')
    if request.method == 'POST':
        year_month = (request.POST.get('year_month') or '').strip()
        if not year_month:
            messages.error(request, '対象年月を指定してください。')
        else:
            _, created = T_ChinaInvoiceMonthClose.objects.get_or_create(
                year_month=year_month, defaults={'closed_by': request.user})
            if created:
                messages.success(request, f'{year_month} を締めました。')
            else:
                messages.error(request, f'{year_month} は既に締め済みです。')
        return redirect('expenses:china_invoice_month_close')

    closed_months = T_ChinaInvoiceMonthClose.objects.order_by('-year_month')
    return render(request, 'expenses/china_invoice_month_close.html', {
        'closed_months': closed_months, 'current': 'china_invoice_month_close',
    })


@login_required
def china_invoice_china_check(request):
    _require_role(request.user, 'china_partner')
    records = T_ChinaInvoice.objects.select_related('cargo_category').order_by('-registered_at')
    params = request.GET
    if params.get('registered_date'):
        records = records.filter(registered_at__date=params['registered_date'])
    if params.get('export_date'):
        records = records.filter(export_date=params['export_date'])
    if params.get('month'):
        records = records.filter(registered_at__date__startswith=params['month'])
    return render(request, 'expenses/china_invoice_china_check.html', {
        'records': records, 'current': 'china_invoice_china_check',
    })


@login_required
@require_POST
def china_invoice_china_check_update(request):
    """チェックした行をまとめて確認済みにする（報告）。"""
    _require_role(request.user, 'china_partner')
    now = timezone.now()

    pks = request.POST.getlist('pks')
    updated = T_ChinaInvoice.objects.filter(pk__in=pks).update(
        china_confirm_status=T_ChinaInvoice.CHINA_STATUS_CONFIRMED,
        china_confirmed_by=request.user, china_confirmed_at=now)

    if updated:
        messages.success(request, f'{updated}件を確認済みにしました。')
    else:
        messages.info(request, '選択された項目がありません。')

    return redirect('expenses:china_invoice_china_check')


@login_required
@require_POST
def china_invoice_china_check_toggle(request, pk):
    """1行分の「差異あり」フラグをAJAXでトグルする。ONで差異あり、OFFで未確認に戻す。"""
    _require_role(request.user, 'china_partner')
    invoice = get_object_or_404(T_ChinaInvoice, pk=pk)
    now = timezone.now()

    if invoice.china_confirm_status == T_ChinaInvoice.CHINA_STATUS_DIFFERENCE:
        invoice.china_confirm_status = T_ChinaInvoice.CHINA_STATUS_UNCONFIRMED
        invoice.china_confirmed_by = None
        invoice.china_confirmed_at = None
    else:
        invoice.china_confirm_status = T_ChinaInvoice.CHINA_STATUS_DIFFERENCE
        invoice.china_confirmed_by = request.user
        invoice.china_confirmed_at = now
    invoice.save(update_fields=['china_confirm_status', 'china_confirmed_by', 'china_confirmed_at'])

    return JsonResponse({
        'china_confirm_status': invoice.china_confirm_status,
        'is_difference': invoice.china_confirm_status == T_ChinaInvoice.CHINA_STATUS_DIFFERENCE,
        'status_display': invoice.get_china_confirm_status_display(),
    })


# 検印欄付き帳票として出力する（経理確認・中国確認の項目は出力しない）
# 「通貨」以降の6列は INVOICE実績報告書（Excel手作業帳票）のK〜P列に相当する。
# 為替レートはDBに持たないため出力後にレートセル（_EXCEL_RATE_CELL）へ手入力する運用とし、
# JPY換算3列はレートセルを参照する数式で自動計算させる。
_EXCEL_HEADERS = [
    '管理番号', 'Invoice No', 'Invoice Total', '輸出日', '貨物\n概要\n区分', '貨物\n概要\n補足',
    '加算調整率', '報告者', '登録日時',
    '通貨', '元値相当（通貨）', '管理費（通貨）', '元値相当（JPY）', '管理費（JPY）', '金額（JPY）',
]
# 帳票レイアウト: 1-2行目=タイトル・検印欄、3行目=為替レート入力欄、4行目=表ヘッダー、5行目〜=データ、最終行=合計
_EXCEL_HEADER_ROW = 4
_EXCEL_DATA_START_ROW = 5
_STAMP_LABELS = ('承認', '確認', '担当')
_STAMP_START_COL = 5  # 検印欄の開始列（E列。E〜Gに3枠並べる）
# 為替レート手入力セル（3行目・「管理費（JPY）」列＝N列。位置を変えたら数式参照も変わるためここで一元管理）
_EXCEL_RATE_CELL = '$N$3'
_EXCEL_DEFAULT_CURRENCY = 'US$'


def _china_invoice_to_excel_row(r, row_idx):
    # 加算調整率(G列)は%値（5.00=5%）のため /100 して使う。JPY換算列は通貨がJPYならレートを掛けない
    rate = _EXCEL_RATE_CELL
    return [
        r.management_no, r.invoice_no, float(r.invoice_total), r.export_date,
        r.cargo_category.content, r.cargo_note, float(r.adjustment_rate_value),
        r.reporter.user_name, r.registered_at.replace(tzinfo=None),
        _EXCEL_DEFAULT_CURRENCY,
        f'=C{row_idx}/(1+G{row_idx}/100)',
        f'=C{row_idx}-K{row_idx}',
        f'=IF({rate}="","",O{row_idx}-N{row_idx})',
        f'=IF({rate}="","",ROUND(IF(J{row_idx}<>"JPY",{rate},1)*L{row_idx},0))',
        f'=IF({rate}="","",ROUND(IF(J{row_idx}<>"JPY",{rate},1)*C{row_idx},0))',
    ]


@login_required
def china_invoice_excel(request):
    _require_role(request.user, *_LIST_ROLES)
    year_month = request.GET.get('year_month')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    records = T_ChinaInvoice.objects.select_related('cargo_category', 'reporter')
    filename = '中国輸出実績_全件.xlsx'
    period_label = '全件'
    if year_month:
        records = records.filter(registered_at__date__startswith=year_month)
        filename = f'中国輸出実績_{year_month.replace("-", "")}.xlsx'
        period_label = year_month
    elif date_from and date_to:
        records = records.filter(registered_at__date__gte=date_from, registered_at__date__lte=date_to)
        filename = f'中国輸出実績_{date_from.replace("-", "")}-{date_to.replace("-", "")}.xlsx'
        period_label = f'{date_from} 〜 {date_to}'
    records = records.order_by('invoice_no')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '中国輸出Invoice'
    # 帳票見映え優先: 表は罫線なし（検印欄の押印枠のみ罫線）＋縞模様で見せる
    ws.sheet_view.showGridLines = False
    n_cols = len(_EXCEL_HEADERS)
    thin = Side(style='thin')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    # タイトル・対象期間（検印欄の左側、A〜D列にマージ）
    title_span = _STAMP_START_COL - 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_span)
    ws.cell(row=1, column=1, value='中国輸出実績報告').font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=title_span)
    ws.cell(row=2, column=1,
            value=f'対象期間: {period_label}　出力日: {datetime.date.today().isoformat()}')

    # 検印欄（E〜G列の3枠: ラベル行＋押印用の空欄行）
    for i, label in enumerate(_STAMP_LABELS):
        col = _STAMP_START_COL + i
        cell = ws.cell(row=1, column=col, value=label)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = box
        ws.cell(row=2, column=col).border = box
    ws.row_dimensions[2].height = 45

    # 為替レート入力欄（_EXCEL_RATE_CELL と同じ位置に配置すること）
    rate_col = coordinate_to_tuple(_EXCEL_RATE_CELL.replace('$', ''))[1]
    label_cell = ws.cell(row=3, column=rate_col - 1, value='為替レート')
    label_cell.alignment = Alignment(horizontal='right', vertical='center')
    rate_cell = ws.cell(row=3, column=rate_col)
    rate_cell.border = box
    rate_cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
    rate_cell.number_format = '#,##0.00'

    # 表ヘッダー
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='495057', end_color='495057', fill_type='solid')
    for col_idx, header in enumerate(_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=_EXCEL_HEADER_ROW, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    # 貨物概要区分・補足の3行折り返しヘッダーが収まる高さにする
    ws.row_dimensions[_EXCEL_HEADER_ROW].height = 48

    stripe_fill = PatternFill(start_color='EDEFF2', end_color='EDEFF2', fill_type='solid')
    amount_col = _EXCEL_HEADERS.index('Invoice Total') + 1
    decimal_cols = {amount_col} | {
        _EXCEL_HEADERS.index(h) + 1 for h in ('元値相当（通貨）', '管理費（通貨）')}
    jpy_cols = {
        _EXCEL_HEADERS.index(h) + 1 for h in ('元値相当（JPY）', '管理費（JPY）', '金額（JPY）')}
    total_amount = Decimal('0')
    record_count = 0
    last_data_row = _EXCEL_DATA_START_ROW - 1
    for row_idx, r in enumerate(records.iterator(), start=_EXCEL_DATA_START_ROW):
        total_amount += r.invoice_total
        record_count += 1
        last_data_row = row_idx
        striped = (row_idx - _EXCEL_DATA_START_ROW) % 2 == 1
        for col_idx, value in enumerate(_china_invoice_to_excel_row(r, row_idx), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if striped:
                cell.fill = stripe_fill
            if col_idx in decimal_cols:
                cell.number_format = '#,##0.00'
            elif col_idx in jpy_cols:
                cell.number_format = '#,##0'

    # 合計行（Invoice Total と換算系列の合計、件数）
    total_row = last_data_row + 1
    total_fill = PatternFill(start_color='D9DEE4', end_color='D9DEE4', fill_type='solid')
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = total_fill
    ws.cell(row=total_row, column=1, value=f'合計（{record_count}件）')
    total_cell = ws.cell(row=total_row, column=amount_col, value=float(total_amount))
    total_cell.number_format = '#,##0.00'
    if record_count:
        for col_idx in sorted(decimal_cols - {amount_col} | jpy_cols):
            letter = get_column_letter(col_idx)
            cell = ws.cell(
                row=total_row, column=col_idx,
                value=f'=SUM({letter}{_EXCEL_DATA_START_ROW}:{letter}{last_data_row})')
            cell.number_format = '#,##0.00' if col_idx in decimal_cols else '#,##0'

    for col_idx, width in enumerate(
            [16, 16, 14, 12, 7, 7, 10, 14, 18, 8, 15, 15, 15, 15, 15], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = f'A{_EXCEL_DATA_START_ROW}'
    ws.auto_filter.ref = (
        f'A{_EXCEL_HEADER_ROW}:{get_column_letter(n_cols)}{max(last_data_row, _EXCEL_HEADER_ROW)}')

    # 印刷設定: A4横・横1ページ収め・左右余白0.5インチ・表ヘッダーを各ページに繰り返し
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f'{_EXCEL_HEADER_ROW}:{_EXCEL_HEADER_ROW}'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    # 日本語ファイル名は RFC 5987 (filename*=UTF-8''...) で出力する。
    # f'filename="{fname}"' 直書きだと Django が非Latin-1ヘッダを RFC 2047 で
    # エンコードし、ブラウザが解釈できず既定名になる（expenses/views.py の
    # データ出力CSVで既に踏んでいる既知の落とし穴と同じ対処）。
    response['Content-Disposition'] = content_disposition_header(as_attachment=True, filename=filename)
    wb.save(response)
    return response
