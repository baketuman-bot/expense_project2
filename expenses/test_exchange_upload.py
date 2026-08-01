"""見出し変換マスタ (M_ExchangeField) と汎用アップロードロジックのテスト"""
from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse

from expenses.models import M_ExchangeField, T_ChinaExport

User = get_user_model()


class MExchangeFieldModelTests(TestCase):
    def test_同一table_nameとupdata_titleの組み合わせは重複登録できない(self):
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                M_ExchangeField.objects.create(
                    table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')

    def test_table_nameが違えば同じupdata_titleを登録できる(self):
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')
        other = M_ExchangeField.objects.create(
            table_name='t_other_table', updata_title='注文番号', up_field_name='order_no')
        self.assertIsNotNone(other.pk)


class MExchangeFieldMasterSettingsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username='master_tester', man_number='9201',
            user_name='マスタ担当', password='pass')

    def test_一覧画面に表示される(self):
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')
        self.client.force_login(self.user)
        res = self.client.get(reverse('expenses:settings_master_list', args=['m_exchange_fields']))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 't_china_export')
        self.assertContains(res, '注文番号')

    def test_新規作成できる(self):
        self.client.force_login(self.user)
        res = self.client.post(
            reverse('expenses:settings_master_create', args=['m_exchange_fields']),
            {'table_name': 't_china_export', 'updata_title': '金額', 'up_field_name': 'amount'})
        self.assertRedirects(
            res, reverse('expenses:settings_master_list', args=['m_exchange_fields']))
        self.assertTrue(
            M_ExchangeField.objects.filter(table_name='t_china_export', updata_title='金額').exists())


from expenses.exchange_upload import get_field_mapping, resolve_model


class GetFieldMappingTests(TestCase):
    def test_table_nameに一致するマッピングのみ辞書で返る(self):
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='金額', up_field_name='amount')
        M_ExchangeField.objects.create(
            table_name='t_other_table', updata_title='注文番号', up_field_name='order_no')
        mapping = get_field_mapping('t_china_export')
        self.assertEqual(mapping, {'注文番号': 'order_no', '金額': 'amount'})

    def test_マッピングが無ければ空辞書(self):
        self.assertEqual(get_field_mapping('t_nonexistent'), {})


class ResolveModelTests(TestCase):
    def test_db_tableが一致するモデルを返す(self):
        self.assertIs(resolve_model('t_china_export'), T_ChinaExport)

    def test_一致するモデルが無ければNone(self):
        self.assertIsNone(resolve_model('t_nonexistent_table'))


import io
from decimal import Decimal
from datetime import date

import openpyxl

from expenses.exchange_upload import parse_excel_rows


def _build_workbook(header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_empty_workbook():
    """見出し行すら持たない完全に空のブックを作る（iter_rowsが1行も返さない）。"""
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ParseExcelRowsTests(TestCase):
    def setUp(self):
        self.mapping = {
            '注文番号': 'order_no',
            '品目名1': 'item_name1',
            '金額': 'amount',
            '購入日': 'purchase_date',
            '無関係の列': 'no_such_field_should_be_ignored_by_mapping_absence',
        }
        # マッピングに存在しない列は mapping.get() で None になるため、
        # このテストでは「マッピングされている4列」のみを対象にする
        del self.mapping['無関係の列']

    def test_正常な行が検証済みデータとして返る(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日', '無視される列'],
            [['ORDER001', 'テスト品目', 1000, '2026-07-01', 'ignore-me']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(len(valid_rows), 1)
        self.assertEqual(valid_rows[0]['order_no'], 'ORDER001')
        self.assertEqual(valid_rows[0]['item_name1'], 'テスト品目')
        self.assertEqual(valid_rows[0]['amount'], Decimal('1000'))
        self.assertEqual(valid_rows[0]['purchase_date'], date(2026, 7, 1))

    def test_スラッシュ区切りの日付も解析できる(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [['ORDER002', 'テスト品目2', 500, '2026/07/02']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(valid_rows[0]['purchase_date'], date(2026, 7, 2))

    def test_カンマ区切りの金額も解析できる(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [['ORDER003', 'テスト品目3', '1,234', '2026-07-03']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(valid_rows[0]['amount'], Decimal('1234'))

    def test_必須項目が空だとエラーになり全件保存されない(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [
                ['ORDER004', 'テスト品目4', 2000, '2026-07-04'],
                ['ORDER005', '', 3000, '2026-07-05'],  # 品目名1が空
            ],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(valid_rows, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0]['row'], 3)  # ヘッダーが1行目、データは2行目起算+1
        self.assertEqual(errors[0]['title'], '品目名1')

    def test_マッピングに無い見出し列は無視される(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日', '未定義列'],
            [['ORDER006', 'テスト品目6', 4000, '2026-07-06', 'なにか']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(len(valid_rows), 1)

    def test_データ行が無ければ空リストが返る(self):
        wb_file = _build_workbook(['注文番号', '品目名1', '金額', '購入日'], [])
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(valid_rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn('ありません', errors[0]['message'])

    def test_マッピング先のフィールド名がモデルに存在しないと設定エラーになる(self):
        bad_mapping = {'注文番号': 'this_field_does_not_exist'}
        wb_file = _build_workbook(['注文番号'], [['ORDER007']])
        valid_rows, errors = parse_excel_rows(wb_file, bad_mapping, T_ChinaExport)
        self.assertEqual(valid_rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn('マスタ設定', errors[0]['message'])

    def test_小数を含むfloatセルが誤差なくDecimalに変換される(self):
        # openpyxl は小数を含む数値セルを Python の float で返す。
        # float をそのまま DecimalField に渡すと2進浮動小数の誤差が展開され
        # decimal_places 検証に落ちてファイル全体が弾かれてしまう。
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [['ORDER008', 'テスト品目8', 1234.56, '2026-07-08']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(len(valid_rows), 1)
        self.assertEqual(valid_rows[0]['amount'], Decimal('1234.56'))

    def test_整数値のfloatセルも小数誤差なく解析される(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [['ORDER009', 'テスト品目9', 1500.0, '2026-07-09']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(valid_rows[0]['amount'], Decimal('1500'))

    def test_小数2桁のfloatセルが複数行あっても全て解析される(self):
        wb_file = _build_workbook(
            ['注文番号', '品目名1', '金額', '購入日'],
            [
                ['ORDER010', 'テスト品目10', 99.99, '2026-07-10'],
                ['ORDER011', 'テスト品目11', 0.1, '2026-07-11'],
                ['ORDER012', 'テスト品目12', 1234.56, '2026-07-12'],
            ],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(errors, [])
        self.assertEqual(
            [r['amount'] for r in valid_rows],
            [Decimal('99.99'), Decimal('0.1'), Decimal('1234.56')])

    def test_マッピングに一致する列が1つも無ければエラーになる(self):
        wb_file = _build_workbook(
            ['未定義列A', '未定義列B'],
            [['なにか', 'なにか2']],
        )
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(valid_rows, [])
        self.assertEqual(len(errors), 1)
        self.assertIn('有効な列が見つかりません', errors[0]['message'])

    def test_見出し行すら無い空のブックはデータなしエラーになる(self):
        wb_file = _build_empty_workbook()
        valid_rows, errors = parse_excel_rows(wb_file, self.mapping, T_ChinaExport)
        self.assertEqual(valid_rows, [])
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0]['message'], 'ファイルにデータがありません')
