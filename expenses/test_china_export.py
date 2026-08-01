"""中国輸出実績報告 (T_ChinaExport) のテスト"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse

from expenses.models import M_UserRole, T_ChinaExport

User = get_user_model()


class TChinaExportModelTests(TestCase):
    def test_品目名1がNullだと保存時にエラー(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                T_ChinaExport.objects.create(item_name1=None, amount=Decimal('100.00'))

    def test_金額がNullだと保存時にエラー(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                T_ChinaExport.objects.create(item_name1='テスト品目', amount=None)

    def test_品目名1と金額以外はNullで保存できる(self):
        record = T_ChinaExport.objects.create(
            item_name1='テスト品目', amount=Decimal('1000.00'))
        self.assertIsNone(record.order_no)
        self.assertIsNone(record.purchase_date)
        self.assertIsNone(record.export_planned_date)
        self.assertIsNone(record.export_date)
        self.assertIsNone(record.invoice_no)
        self.assertIsNone(record.updated_by)


class ChinaExportListViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.export_user = User.objects.create_user(
            username='export_tester', man_number='9101',
            user_name='輸出担当', password='pass')
        M_UserRole.objects.create(man_number=cls.export_user, role='export')

        cls.other_user = User.objects.create_user(
            username='other_tester', man_number='9102',
            user_name='権限なし', password='pass')

        cls.admin_user = User.objects.create_user(
            username='admin_tester', man_number='9107',
            user_name='管理者', password='pass')
        M_UserRole.objects.create(man_number=cls.admin_user, role='admin')

        cls.unexported = T_ChinaExport.objects.create(
            item_name1='未輸出品', amount=Decimal('5000.00'),
            purchase_date=date(2026, 6, 1),
            supplier_name='上海サプライヤー', order_staff_name='山田太郎',
            account_name='仕入高', burden_bumon_name='営業部')
        cls.exported = T_ChinaExport.objects.create(
            item_name1='輸出済品', amount=Decimal('3000.00'),
            purchase_date=date(2026, 5, 1),
            export_date=date(2026, 6, 15))

    def test_exportロールを持たないユーザーは403(self):
        self.client.force_login(self.other_user)
        res = self.client.get(reverse('expenses:china_export_list'))
        self.assertEqual(res.status_code, 403)

    def test_adminロールを持っていればexportロールがなくても閲覧できる(self):
        self.client.force_login(self.admin_user)
        res = self.client.get(reverse('expenses:china_export_list'))
        self.assertEqual(res.status_code, 200)

    def test_経理入力項目がテンプレートに表示される(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list'))
        self.assertContains(res, '上海サプライヤー')
        self.assertContains(res, '山田太郎')
        self.assertContains(res, '仕入高')
        self.assertContains(res, '営業部')

    def test_デフォルトは未輸出のみ表示(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, '未輸出品')
        self.assertNotContains(res, '輸出済品')

    def test_show_allで全件表示(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?show=all')
        self.assertContains(res, '未輸出品')
        self.assertContains(res, '輸出済品')

    def test_購入日昇順で並ぶ(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?show=all')
        records = list(res.context['records'])
        self.assertEqual(
            [r.pk for r in records],
            [self.exported.pk, self.unexported.pk])

    def test_sortパラメータで指定フィールド昇順に並ぶ(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?show=all&sort=amount')
        records = list(res.context['records'])
        self.assertEqual(
            [r.pk for r in records],
            [self.exported.pk, self.unexported.pk])  # amount: 3000(exported) < 5000(unexported)

    def test_sortパラメータの先頭にマイナスを付けると降順に並ぶ(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?show=all&sort=-amount')
        records = list(res.context['records'])
        self.assertEqual(
            [r.pk for r in records],
            [self.unexported.pk, self.exported.pk])  # amount: 5000(unexported) > 3000(exported)

    def test_不正なsortパラメータはデフォルト順にフォールバックする(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?show=all&sort=__class__')
        self.assertEqual(res.status_code, 200)
        records = list(res.context['records'])
        self.assertEqual(
            [r.pk for r in records],
            [self.exported.pk, self.unexported.pk])  # purchase_date昇順(デフォルト)

    def test_ソート中の列見出しにアクティブ状態が反映される(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list') + '?sort=amount')
        self.assertTrue(res.context['sort_links']['amount']['active'])
        self.assertEqual(res.context['sort_links']['amount']['arrow'], '▲')
        self.assertFalse(res.context['sort_links']['order_no']['active'])

    def test_exportロールを持たないユーザーはExcel出力不可(self):
        self.client.force_login(self.other_user)
        res = self.client.get(reverse('expenses:china_export_excel'))
        self.assertEqual(res.status_code, 403)

    def test_Excel出力は未輸出のみデフォルト表示され経理項目を含む(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_excel'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        values = [cell.value for row in wb.active.iter_rows() for cell in row]
        self.assertIn('注文番号', values)
        self.assertIn('上海サプライヤー', values)
        self.assertNotIn('輸出済品', values)

    def test_Excel出力はshow_allで全件含む(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_excel') + '?show=all')
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        values = [cell.value for row in wb.active.iter_rows() for cell in row]
        self.assertIn('未輸出品', values)
        self.assertIn('輸出済品', values)

    def test_Excel出力は金額列に数値とカンマ書式が設定される(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_excel'))
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        amount_col = header.index('金額') + 1
        data_cell = ws.cell(row=2, column=amount_col)
        self.assertEqual(data_cell.value, 5000.0)
        self.assertEqual(data_cell.number_format, '#,##0')

    def test_一覧画面にアップロードボタンがある(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_list'))
        self.assertContains(res, reverse('expenses:china_export_upload'))


class ChinaExportBulkUpdateViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.export_user = User.objects.create_user(
            username='export_tester2', man_number='9103',
            user_name='輸出担当2', password='pass')
        M_UserRole.objects.create(man_number=cls.export_user, role='export')
        cls.other_user = User.objects.create_user(
            username='other_tester2', man_number='9104',
            user_name='権限なし2', password='pass')
        cls.admin_user = User.objects.create_user(
            username='admin_tester2', man_number='9108',
            user_name='管理者2', password='pass')
        M_UserRole.objects.create(man_number=cls.admin_user, role='admin')
        cls.record = T_ChinaExport.objects.create(
            order_no='ORDER0001', item_name1='対象品目', amount=Decimal('1234.00'))
        cls.record2 = T_ChinaExport.objects.create(
            order_no='ORDER0002', item_name1='対象品目2', amount=Decimal('5678.00'))

    def test_exportロールを持たないユーザーは一括更新不可(self):
        self.client.force_login(self.other_user)
        res = self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk],
                'export_planned_date_' + str(self.record.pk): '2026-08-01',
                'export_date_' + str(self.record.pk): '',
                'invoice_no_' + str(self.record.pk): 'INV-001',
            })
        self.assertEqual(res.status_code, 403)

    def test_adminロールを持っていればexportロールがなくても一括更新できる(self):
        self.client.force_login(self.admin_user)
        res = self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk],
                'export_planned_date_' + str(self.record.pk): '2026-08-01',
                'export_date_' + str(self.record.pk): '',
                'invoice_no_' + str(self.record.pk): 'INV-001',
            })
        self.assertRedirects(res, reverse('expenses:china_export_list'))
        self.record.refresh_from_db()
        self.assertEqual(self.record.export_planned_date, date(2026, 8, 1))

    def test_複数行をまとめて一括保存できる(self):
        original_updated_at = self.record.updated_at
        self.client.force_login(self.export_user)
        res = self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk, self.record2.pk],
                'export_planned_date_' + str(self.record.pk): '2026-08-01',
                'export_date_' + str(self.record.pk): '2026-08-10',
                'invoice_no_' + str(self.record.pk): 'INV-001',
                'export_planned_date_' + str(self.record2.pk): '2026-09-01',
                'export_date_' + str(self.record2.pk): '',
                'invoice_no_' + str(self.record2.pk): 'INV-002',
            })
        self.assertRedirects(res, reverse('expenses:china_export_list'))
        self.record.refresh_from_db()
        self.record2.refresh_from_db()
        self.assertEqual(self.record.export_planned_date, date(2026, 8, 1))
        self.assertEqual(self.record.export_date, date(2026, 8, 10))
        self.assertEqual(self.record.invoice_no, 'INV-001')
        self.assertEqual(self.record.updated_by, self.export_user)
        self.assertIsNotNone(self.record.updated_at)
        self.assertNotEqual(self.record.updated_at, original_updated_at)
        self.assertEqual(self.record2.export_planned_date, date(2026, 9, 1))
        self.assertIsNone(self.record2.export_date)
        self.assertEqual(self.record2.invoice_no, 'INV-002')

    def test_経理入力項目はPOSTに含めても更新されない(self):
        self.client.force_login(self.export_user)
        self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk],
                'export_planned_date_' + str(self.record.pk): '',
                'export_date_' + str(self.record.pk): '',
                'invoice_no_' + str(self.record.pk): '',
                'order_no': 'HACKED', 'amount': '999999.00',
            })
        self.record.refresh_from_db()
        self.assertEqual(self.record.order_no, 'ORDER0001')
        self.assertEqual(self.record.amount, Decimal('1234.00'))

    def test_pksに含まれないレコードは更新されない(self):
        self.client.force_login(self.export_user)
        self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk],
                'export_planned_date_' + str(self.record.pk): '2026-08-01',
                'export_date_' + str(self.record.pk): '',
                'invoice_no_' + str(self.record.pk): '',
                # record2 用のフィールドを送っても pks に含めなければ更新されない
                'export_planned_date_' + str(self.record2.pk): '2026-09-01',
            })
        self.record2.refresh_from_db()
        self.assertIsNone(self.record2.export_planned_date)

    def test_show_allを維持したままリダイレクトされる(self):
        self.client.force_login(self.export_user)
        res = self.client.post(
            reverse('expenses:china_export_bulk_update'),
            {
                'pks': [self.record.pk],
                'export_planned_date_' + str(self.record.pk): '',
                'export_date_' + str(self.record.pk): '',
                'invoice_no_' + str(self.record.pk): '',
                'show': 'all',
            })
        self.assertRedirects(res, reverse('expenses:china_export_list') + '?show=all')


class ChinaExportSidebarTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.export_user = User.objects.create_user(
            username='export_tester3', man_number='9105',
            user_name='輸出担当3', password='pass')
        M_UserRole.objects.create(man_number=cls.export_user, role='export')
        cls.other_user = User.objects.create_user(
            username='other_tester3', man_number='9106',
            user_name='権限なし3', password='pass')
        cls.admin_user = User.objects.create_user(
            username='admin_tester3', man_number='9109',
            user_name='管理者3', password='pass')
        M_UserRole.objects.create(man_number=cls.admin_user, role='admin')

    def test_exportロール保持者はサイドバーにメニューが出る(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, '中国輸出実績報告')

    def test_adminロール保持者もサイドバーにメニューが出る(self):
        self.client.force_login(self.admin_user)
        res = self.client.get(reverse('expenses:home'))
        self.assertContains(res, '中国輸出実績報告')

    def test_exportロールもadminロールもないユーザーには出ない(self):
        self.client.force_login(self.other_user)
        res = self.client.get(reverse('expenses:home'))
        self.assertNotContains(res, '中国輸出実績報告')


import json
from datetime import datetime

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from expenses.models import M_ExchangeField
from expenses.views_china_export import _deserialize_staged_value, _serialize_staged_value


def _build_china_export_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['注文番号', '品目名1', '金額', '購入日'])
    ws.append(['UP0001', 'アップロード品目', 1500, '2026-07-10'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return SimpleUploadedFile(
        'upload.xlsx', buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class ChinaExportUploadViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.export_user = User.objects.create_user(
            username='export_tester4', man_number='9110',
            user_name='輸出担当4', password='pass')
        M_UserRole.objects.create(man_number=cls.export_user, role='export')
        cls.other_user = User.objects.create_user(
            username='other_tester4', man_number='9111',
            user_name='権限なし4', password='pass')
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='注文番号', up_field_name='order_no')
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='品目名1', up_field_name='item_name1')
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='金額', up_field_name='amount')
        M_ExchangeField.objects.create(
            table_name='t_china_export', updata_title='購入日', up_field_name='purchase_date')

    def test_exportロールを持たないユーザーは403(self):
        self.client.force_login(self.other_user)
        res = self.client.get(reverse('expenses:china_export_upload'))
        self.assertEqual(res.status_code, 403)

    def test_GET画面表示でセッションがクリアされる(self):
        self.client.force_login(self.export_user)
        session = self.client.session
        session['china_export_upload_staged'] = [{'dummy': 'x'}]
        session.save()
        self.client.get(reverse('expenses:china_export_upload'))
        self.assertNotIn('china_export_upload_staged', self.client.session)

    def test_xlsx以外の拡張子はエラー表示(self):
        self.client.force_login(self.export_user)
        bad_file = SimpleUploadedFile('upload.csv', b'a,b,c', content_type='text/csv')
        res = self.client.post(reverse('expenses:china_export_upload'), {'excel_file': bad_file})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, '.xlsxのみ')

    def test_マッピング未登録の場合は案内表示(self):
        M_ExchangeField.objects.all().delete()
        self.client.force_login(self.export_user)
        res = self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, '見出し変換マスタが未設定です')

    def test_正常なファイルはプレビュー表示され保存はされない(self):
        self.client.force_login(self.export_user)
        res = self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'アップロード品目')
        self.assertEqual(res.context['preview_count'], 1)
        self.assertFalse(T_ChinaExport.objects.filter(order_no='UP0001').exists())

    def test_プレビュー画面の確定フォームに二重送信防止属性がある(self):
        self.client.force_login(self.export_user)
        res = self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.assertContains(res, 'data-confirm-form')

    def test_エラーがある場合はエラー一覧が表示されデータは保存されない(self):
        self.client.force_login(self.export_user)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['注文番号', '品目名1', '金額', '購入日'])
        ws.append(['UP0002', '', 1500, '2026-07-10'])  # 品目名1が空
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        bad_file = SimpleUploadedFile(
            'upload2.xlsx', buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        res = self.client.post(reverse('expenses:china_export_upload'), {'excel_file': bad_file})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, '品目名1')
        self.assertFalse(T_ChinaExport.objects.filter(order_no='UP0002').exists())

    def test_プレビューなしでconfirmにPOSTしても保存されず案内される(self):
        self.client.force_login(self.export_user)
        res = self.client.post(reverse('expenses:china_export_upload_confirm'), follow=True)
        self.assertRedirects(res, reverse('expenses:china_export_upload'))
        self.assertEqual(T_ChinaExport.objects.filter(order_no='UP0001').count(), 0)
        # リダイレクト先でメッセージが実際に表示され、Bootstrapのクラスも正しいこと
        self.assertContains(res, 'アップロードするデータがありません')
        self.assertContains(res, 'alert-danger')
        self.assertNotContains(res, 'alert-error')

    def test_アップロード画面はサイドバーの中国輸出実績報告を選択状態にする(self):
        self.client.force_login(self.export_user)
        res = self.client.get(reverse('expenses:china_export_upload'))
        self.assertEqual(res.context['current'], 'china_export_list')
        res = self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.assertEqual(res.context['current'], 'china_export_list')

    def test_アップロード失敗時は以前ステージしたデータが破棄される(self):
        self.client.force_login(self.export_user)
        # 1回目: 正常なファイル → セッションにステージされる
        self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.assertIn('china_export_upload_staged', self.client.session)
        # 2回目: 不正なファイル → 以前のステージデータは残らない
        bad_file = SimpleUploadedFile('upload.csv', b'a,b,c', content_type='text/csv')
        res = self.client.post(reverse('expenses:china_export_upload'), {'excel_file': bad_file})
        self.assertEqual(res.status_code, 200)
        self.assertNotIn('china_export_upload_staged', self.client.session)
        # 確定しても何も保存されない
        self.client.post(reverse('expenses:china_export_upload_confirm'))
        self.assertEqual(T_ChinaExport.objects.filter(order_no='UP0001').count(), 0)

    def test_上限件数を超えるファイルはエラー表示されステージされない(self):
        from expenses.views_china_export import _UPLOAD_MAX_ROWS

        self.client.force_login(self.export_user)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['注文番号', '品目名1', '金額', '購入日'])
        for i in range(_UPLOAD_MAX_ROWS + 1):
            ws.append([f'UPMAX{i:05d}', '大量品目', 100, '2026-07-10'])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        big_file = SimpleUploadedFile(
            'upload_big.xlsx', buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        res = self.client.post(reverse('expenses:china_export_upload'), {'excel_file': big_file})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, f'一度にアップロードできるのは{_UPLOAD_MAX_ROWS}件までです')
        self.assertNotIn('china_export_upload_staged', self.client.session)
        self.assertEqual(T_ChinaExport.objects.filter(item_name1='大量品目').count(), 0)

    def test_xlsx拡張子でも中身が壊れていればエラー表示される(self):
        self.client.force_login(self.export_user)
        broken = SimpleUploadedFile(
            'upload.xlsx', b'this is not a real xlsx file',
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        with self.assertLogs('expenses.views_china_export', level='ERROR'):
            res = self.client.post(reverse('expenses:china_export_upload'), {'excel_file': broken})
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'ファイルの読み込みに失敗しました')
        self.assertNotIn('china_export_upload_staged', self.client.session)

    def test_プレビュー後に確定するとDBへ保存される(self):
        self.client.force_login(self.export_user)
        self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        res = self.client.post(reverse('expenses:china_export_upload_confirm'))
        self.assertRedirects(res, reverse('expenses:china_export_list'))
        record = T_ChinaExport.objects.get(order_no='UP0001')
        self.assertEqual(record.item_name1, 'アップロード品目')
        self.assertEqual(record.amount, Decimal('1500'))
        self.assertEqual(record.purchase_date, date(2026, 7, 10))
        self.assertNotIn('china_export_upload_staged', self.client.session)

    def test_確定後は同じ内容を再度確定しても増えない(self):
        self.client.force_login(self.export_user)
        self.client.post(
            reverse('expenses:china_export_upload'), {'excel_file': _build_china_export_workbook()})
        self.client.post(reverse('expenses:china_export_upload_confirm'))
        self.client.post(reverse('expenses:china_export_upload_confirm'))  # セッションは既に消えている
        self.assertEqual(T_ChinaExport.objects.filter(order_no='UP0001').count(), 1)

    def test_exportロールを持たないユーザーは確定不可(self):
        self.client.force_login(self.other_user)
        res = self.client.post(reverse('expenses:china_export_upload_confirm'))
        self.assertEqual(res.status_code, 403)


class StagedValueSerializationTests(SimpleTestCase):
    """セッション保存用の値変換 (_serialize_staged_value / _deserialize_staged_value)"""

    def _roundtrip(self, value):
        return _deserialize_staged_value(json.loads(json.dumps(_serialize_staged_value(value))))

    def test_Decimalが往復できる(self):
        self.assertEqual(self._roundtrip(Decimal('1234.56')), Decimal('1234.56'))

    def test_dateが往復できる(self):
        self.assertEqual(self._roundtrip(date(2026, 7, 10)), date(2026, 7, 10))

    def test_datetimeがdateとして壊れず往復できる(self):
        # datetime は date のサブクラスなので、date 分岐が先だと復元に失敗する
        value = datetime(2026, 7, 10, 13, 45, 30)
        restored = self._roundtrip(value)
        self.assertIsInstance(restored, datetime)
        self.assertEqual(restored, value)

    def test_文字列やNoneはそのまま(self):
        self.assertEqual(self._roundtrip('ORDER001'), 'ORDER001')
        self.assertIsNone(self._roundtrip(None))
